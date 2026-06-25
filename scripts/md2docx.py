import re, sys
from docx import Document
from docx.shared import Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

def set_rtl(par):
    p = par._p.get_or_add_pPr()
    bidi = OxmlElement('w:bidi'); p.append(bidi)
    par.alignment = WD_ALIGN_PARAGRAPH.RIGHT

def rtl_run(run):
    rpr = run._r.get_or_add_rPr()
    rtl = OxmlElement('w:rtl'); rpr.append(rtl)

def add_inline(par, text, mono=False):
    # bold **...**
    parts = re.split(r'(\*\*.+?\*\*)', text)
    for pt in parts:
        if not pt: continue
        if pt.startswith('**') and pt.endswith('**'):
            r = par.add_run(pt[2:-2]); r.bold = True
        else:
            r = par.add_run(pt)
        if mono: r.font.name = 'Consolas'; r.font.size = Pt(9)
        rtl_run(r)

def shade(cell, color):
    tcPr = cell._tc.get_or_add_tcPr()
    sh = OxmlElement('w:shd'); sh.set(qn('w:val'),'clear'); sh.set(qn('w:fill'),color); tcPr.append(sh)

def convert(md_path, docx_path):
    doc = Document()
    style = doc.styles['Normal']; style.font.name='Arial'; style.font.size=Pt(11)
    lines = open(md_path, encoding='utf-8').read().split('\n')
    i=0; n=len(lines)
    while i < n:
        line = lines[i]
        # code fence
        if line.strip().startswith('```'):
            i+=1; buf=[]
            while i<n and not lines[i].strip().startswith('```'):
                buf.append(lines[i]); i+=1
            i+=1
            p=doc.add_paragraph()
            r=p.add_run('\n'.join(buf)); r.font.name='Consolas'; r.font.size=Pt(9)
            continue
        # table
        if line.strip().startswith('|') and i+1<n and re.match(r'^\s*\|[\s:|-]+\|\s*$', lines[i+1]):
            header=[c.strip() for c in line.strip().strip('|').split('|')]
            i+=2; rows=[]
            while i<n and lines[i].strip().startswith('|'):
                rows.append([c.strip() for c in lines[i].strip().strip('|').split('|')]); i+=1
            t=doc.add_table(rows=1, cols=len(header)); t.style='Light Grid Accent 1'
            for j,h in enumerate(header):
                c=t.rows[0].cells[j]; c.paragraphs[0].clear(); add_inline(c.paragraphs[0], h); set_rtl(c.paragraphs[0]); shade(c,'D9E2F3')
            for row in rows:
                cells=t.add_row().cells
                for j in range(len(header)):
                    txt=row[j] if j<len(row) else ''
                    cells[j].paragraphs[0].clear(); add_inline(cells[j].paragraphs[0], txt); set_rtl(cells[j].paragraphs[0])
            doc.add_paragraph()
            continue
        # headings
        m=re.match(r'^(#{1,4})\s+(.*)', line)
        if m:
            lvl=len(m.group(1)); p=doc.add_heading(level=min(lvl,4))
            add_inline(p, m.group(2)); set_rtl(p); 
            i+=1; continue
        if line.strip().startswith('> '):
            p=doc.add_paragraph(); add_inline(p, line.strip()[2:]); 
            for r in p.runs: r.italic=True
            set_rtl(p); i+=1; continue
        if re.match(r'^\s*[-*]\s+', line):
            p=doc.add_paragraph(style='List Bullet'); add_inline(p, re.sub(r'^\s*[-*]\s+','',line)); set_rtl(p); i+=1; continue
        if re.match(r'^\s*\d+\.\s+', line):
            p=doc.add_paragraph(style='List Number'); add_inline(p, re.sub(r'^\s*\d+\.\s+','',line)); set_rtl(p); i+=1; continue
        if line.strip()=='---':
            doc.add_paragraph('─'*30); i+=1; continue
        if line.strip()=='':
            i+=1; continue
        p=doc.add_paragraph(); add_inline(p, line); set_rtl(p); i+=1
    doc.save(docx_path)
    print('saved', docx_path)

convert(sys.argv[1], sys.argv[2])
